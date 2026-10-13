"""Service Reports: builder Excel (openpyxl) & PDF (reportlab) di memori.

Semua fungsi builder murni — menerima data yang sudah dihitung/diambil dan
mengembalikan `io.BytesIO` yang sudah di-`seek(0)` siap dikirim sebagai
`StreamingResponse`. Tidak ada perhitungan ulang: angka RAB/EVM diambil dari
service modul `boq_rab` & `progress`.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from ...models import Inspection, Project
from ...schemas import ProgressSummaryResponse, RABResponse

# --- Gaya umum -------------------------------------------------------------

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_TITLE_FONT = Font(bold=True, size=14)
_SUBTITLE_FONT = Font(bold=True, size=11)
_TOTAL_FONT = Font(bold=True)
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_NUM_CURRENCY = "#,##0"
_NUM_VOLUME = "#,##0.000"
_NUM_PERCENT = "0.00"

# Kolom tabel RAB: (header, atribut RABLineResponse, format angka)
_RAB_COLUMNS: list[tuple[str, str, str | None]] = [
    ("Kode", "code", None),
    ("Uraian", "description", None),
    ("Satuan", "unit", None),
    ("Vol. Kontrak", "contract_volume", _NUM_VOLUME),
    ("Δ Vol.", "delta_volume", _NUM_VOLUME),
    ("Vol. Disesuaikan", "adjusted_volume", _NUM_VOLUME),
    ("Harga Satuan", "unit_price", _NUM_CURRENCY),
    ("Jumlah Kontrak", "contract_amount", _NUM_CURRENCY),
    ("Jumlah Disesuaikan", "adjusted_amount", _NUM_CURRENCY),
    ("Δ Jumlah", "delta_amount", _NUM_CURRENCY),
]


def _enum_value(value):
    """Ambil `.value` bila enum, jika tidak kembalikan apa adanya."""
    return getattr(value, "value", value)


def _save(wb: Workbook) -> io.BytesIO:
    """Simpan workbook ke buffer memori dan kembalikan di posisi awal."""
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _write_header(ws: Worksheet, row: int, headers: list[str]) -> None:
    """Tulis baris header tabel dengan gaya seragam."""
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER


# --- RAB -------------------------------------------------------------------


def build_rab_xlsx(rab: RABResponse) -> io.BytesIO:
    """Workbook Excel "RAB": header proyek, tabel baris RAB, baris total."""
    wb = Workbook()
    ws = wb.active
    ws.title = "RAB"

    ws["A1"] = "Rencana Anggaran Biaya"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Proyek: {rab.project_name}"
    ws["A2"].font = _SUBTITLE_FONT

    header_row = 4
    _write_header(ws, header_row, [c[0] for c in _RAB_COLUMNS])

    row = header_row + 1
    for line in rab.lines:
        for col, (_, attr, num_fmt) in enumerate(_RAB_COLUMNS, start=1):
            cell = ws.cell(row=row, column=col, value=getattr(line, attr))
            cell.border = _BORDER
            if num_fmt:
                cell.number_format = num_fmt
        row += 1

    # Baris total: kolom Kode diisi label, jumlah ditempel di kolom yang sesuai.
    total_row = row
    ws.cell(row=total_row, column=1, value="TOTAL").font = _TOTAL_FONT
    totals = {
        8: rab.total_contract,
        9: rab.total_adjusted,
        10: rab.total_delta,
    }
    for col, value in totals.items():
        cell = ws.cell(row=total_row, column=col, value=value)
        cell.font = _TOTAL_FONT
        cell.number_format = _NUM_CURRENCY
        cell.border = _BORDER

    widths = [12, 40, 10, 14, 12, 16, 16, 18, 18, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=idx).column_letter].width = width

    return _save(wb)


# --- EVM -------------------------------------------------------------------


def build_evm_xlsx(summary: ProgressSummaryResponse) -> io.BytesIO:
    """Workbook Excel "EVM": blok metrik + tabel periode."""
    wb = Workbook()
    ws = wb.active
    ws.title = "EVM"

    ws["A1"] = "Analisis Nilai Hasil (EVM)"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Proyek: {summary.project_name}"
    ws["A2"].font = _SUBTITLE_FONT

    metrics: list[tuple[str, float, str | None]] = [
        ("BAC (Budget at Completion)", summary.bac, _NUM_CURRENCY),
        ("PV (Planned Value)", summary.pv, _NUM_CURRENCY),
        ("EV (Earned Value)", summary.ev, _NUM_CURRENCY),
        ("AC (Actual Cost)", summary.ac, _NUM_CURRENCY),
        ("SV (Schedule Variance)", summary.sv, _NUM_CURRENCY),
        ("CV (Cost Variance)", summary.cv, _NUM_CURRENCY),
        ("SPI (Schedule Performance Index)", summary.spi, "0.0000"),
        ("CPI (Cost Performance Index)", summary.cpi, "0.0000"),
        ("EAC (Estimate at Completion)", summary.eac, _NUM_CURRENCY),
        ("ETC (Estimate to Complete)", summary.etc, _NUM_CURRENCY),
        ("VAC (Variance at Completion)", summary.vac, _NUM_CURRENCY),
    ]

    metrics_row = 4
    _write_header(ws, metrics_row, ["Metrik", "Nilai"])
    row = metrics_row + 1
    for label, value, num_fmt in metrics:
        ws.cell(row=row, column=1, value=label).border = _BORDER
        cell = ws.cell(row=row, column=2, value=value)
        cell.border = _BORDER
        if num_fmt:
            cell.number_format = num_fmt
        row += 1

    periods_row = row + 1
    _write_header(
        ws,
        periods_row,
        ["Periode", "Rencana (%)", "Realisasi (%)", "Nilai Rencana", "Nilai Hasil", "Biaya Aktual"],
    )
    prow = periods_row + 1
    for point in summary.periods:
        values = [
            point.period,
            point.planned_pct,
            point.actual_pct,
            point.planned_value,
            point.earned_value,
            point.actual_cost,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=prow, column=col, value=value)
            cell.border = _BORDER
            if col in (2, 3):
                cell.number_format = _NUM_PERCENT
            elif col >= 4:
                cell.number_format = _NUM_CURRENCY
        prow += 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 18

    return _save(wb)


# --- QC --------------------------------------------------------------------


def build_qc_xlsx(project_name: str, inspections: list[Inspection]) -> io.BytesIO:
    """Workbook Excel "QC": satu baris per butir pemeriksaan inspeksi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "QC"

    ws["A1"] = "Laporan Inspeksi Mutu (QC)"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = f"Proyek: {project_name}"
    ws["A2"].font = _SUBTITLE_FONT

    headers = [
        "No. Inspeksi",
        "Judul",
        "Kategori",
        "Status",
        "Butir Pemeriksaan",
        "Standar",
        "Hasil",
        "Catatan Butir",
    ]
    header_row = 4
    _write_header(ws, header_row, headers)

    row = header_row + 1
    for inspection in inspections:
        for item in inspection.items:
            values = [
                inspection.inspection_number,
                inspection.title,
                _enum_value(inspection.category),
                _enum_value(inspection.status),
                item.description,
                item.standard,
                _enum_value(item.result),
                item.note,
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = _BORDER
            row += 1

    widths = [14, 30, 12, 12, 40, 16, 10, 30]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=idx).column_letter].width = width

    return _save(wb)


# --- RAB PDF ---------------------------------------------------------------


def build_rab_pdf(rab: RABResponse) -> io.BytesIO:
    """PDF "Rencana Anggaran Biaya": tabel baris RAB + ringkasan total."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Rencana Anggaran Biaya",
    )
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Rencana Anggaran Biaya", styles["Title"]),
        Paragraph(f"Proyek: {rab.project_name}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]

    header = ["Kode", "Uraian", "Satuan", "Vol. Kontrak", "Δ Vol.", "Vol. Disesuaikan",
              "Harga Satuan", "Jumlah Kontrak", "Jumlah Disesuaikan", "Δ Jumlah"]
    data: list[list[str]] = [header]
    for line in rab.lines:
        data.append([
            line.code,
            Paragraph(line.description, styles["Normal"]),
            line.unit,
            f"{line.contract_volume:,.3f}",
            f"{line.delta_volume:,.3f}",
            f"{line.adjusted_volume:,.3f}",
            f"{line.unit_price:,.0f}",
            f"{line.contract_amount:,.0f}",
            f"{line.adjusted_amount:,.0f}",
            f"{line.delta_amount:,.0f}",
        ])

    table = Table(data, repeatRows=1, colWidths=[
        16 * mm, 62 * mm, 14 * mm, 20 * mm, 18 * mm, 24 * mm, 22 * mm, 24 * mm, 26 * mm, 20 * mm,
    ])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
    ]))
    elements.append(table)

    elements.append(Spacer(1, 6 * mm))
    elements.append(Paragraph("Rincian Total", styles["Heading3"]))
    totals = Table([
        ["Total Kontrak", f"Rp {rab.total_contract:,.0f}"],
        ["Total Disesuaikan", f"Rp {rab.total_adjusted:,.0f}"],
        ["Δ Total", f"Rp {rab.total_delta:,.0f}"],
    ], colWidths=[50 * mm, 50 * mm])
    totals.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
    ]))
    elements.append(totals)

    doc.build(elements)
    buf.seek(0)
    return buf


# --- Pengambilan data ------------------------------------------------------


def fetch_qc_data(db: Session, project_id: int) -> tuple[str, list[Inspection]]:
    """Ambil nama proyek + seluruh inspeksi (beserta butir) untuk laporan QC.

    Raises:
        ValueError: bila proyek tidak ditemukan (dipetakan router ke HTTP 404).
    """
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise ValueError("Project not found")

    inspections = (
        db.query(Inspection)
        .filter(Inspection.project_id == project_id)
        .order_by(Inspection.id.asc())
        .all()
    )
    return project.name, inspections
